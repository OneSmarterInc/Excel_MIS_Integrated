"""Pull the pictures out of an uploaded book and hang them on the right section.

A textbook explains as much with its screenshots as with its sentences, and until now
only the sentences survived the upload. This reads the same file a second time, takes
every picture out of it, remembers the text that sat around each one, and works out
which chapter or module that text belongs to. The picture is then shown beside the
explanation for that section.

Nothing here is required for an upload to succeed. Every reader is wrapped, and a file
whose pictures cannot be read still produces exactly the chapters and modules it did
before. Pillow is used when it is installed and skipped when it is not.
"""
from __future__ import annotations

import hashlib
import io
import re

# A picture smaller than this is a bullet, a rule or a logo, not a figure worth showing.
MIN_BYTES = 6 * 1024
MIN_PIXELS = 110
# A picture that turns up on page after page is a running header or a watermark.
MAX_REPEATS = 3
MAX_FIGURES = 120
CONTEXT_CHARS = 700

EXTENSIONS = {
    "jpeg": ".jpg", "jpg": ".jpg", "png": ".png", "gif": ".gif",
    "bmp": ".bmp", "tiff": ".tif", "webp": ".webp", "jp2": ".jp2", "emf": ".emf",
    "wmf": ".wmf",
}
# The app draws these. Anything else is stored but will not render on a phone, so it is
# left out rather than shown as a broken box.
SHOWABLE = {".jpg", ".png", ".gif", ".bmp", ".webp"}


def _extension(name: str, content_type: str = "") -> str:
    lowered = (name or "").lower()
    for key, suffix in EXTENSIONS.items():
        if lowered.endswith(f".{key}"):
            return suffix
    guess = (content_type or "").lower().split("/")[-1]
    return EXTENSIONS.get(guess, ".png")


def _dimensions(data: bytes) -> tuple[int, int]:
    """Width and height when Pillow is available, and (0, 0) when it is not."""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as picture:
            return int(picture.width), int(picture.height)
    except Exception:
        return 0, 0


def _worth_keeping(data: bytes, width: int, height: int) -> bool:
    """Is this a figure, or is it furniture?

    When Pillow is installed the size on screen decides it, because a clean diagram can
    compress to almost nothing and still be the most useful thing on the page. Without
    Pillow there is nothing to go on but the weight of the file.
    """
    if width and height:
        if width < MIN_PIXELS or height < MIN_PIXELS:
            return False
        # A very long thin strip is a divider rule rather than a figure.
        if width / max(height, 1) > 12 or height / max(width, 1) > 12:
            return False
        return True
    return len(data) >= MIN_BYTES


def _tidy(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


# --------------------------------------------------------------------------- readers

def _pdf_page_images(page) -> list[tuple[bytes, str]]:
    """Every picture on one PDF page, as (bytes, file name).

    pypdf's own image reader needs Pillow, so when Pillow is missing the compressed
    streams are lifted straight out of the page's resources instead. That covers the
    JPEG screenshots that make up most of a textbook.
    """
    found: list[tuple[bytes, str]] = []
    try:
        for index, image in enumerate(page.images):
            data = getattr(image, "data", None)
            if data:
                found.append((data, getattr(image, "name", f"image{index}.png")))
        if found:
            return found
    except Exception:
        pass

    try:
        resources = page.get("/Resources")
        if resources is None:
            return found
        xobjects = resources.get_object().get("/XObject")
        if xobjects is None:
            return found
        xobjects = xobjects.get_object()
        for index, key in enumerate(xobjects.keys()):
            try:
                obj = xobjects[key].get_object()
                if obj.get("/Subtype") != "/Image":
                    continue
                filters = obj.get("/Filter")
                names = [str(filters)] if not isinstance(filters, list) else [str(f) for f in filters]
                if "/DCTDecode" in names:
                    suffix = ".jpg"
                elif "/JPXDecode" in names:
                    suffix = ".jp2"
                else:
                    continue                     # needs decompressing, which needs Pillow
                found.append((obj.get_data(), f"{str(key).strip('/')}{index}{suffix}"))
            except Exception:
                continue
    except Exception:
        pass
    return found


def read_pdf_figures(path: str) -> list[dict]:
    from pypdf import PdfReader

    figures: list[dict] = []
    reader = PdfReader(path)
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            page_text = _tidy(page.extract_text() or "")
        except Exception:
            page_text = ""
        for order, (data, name) in enumerate(_pdf_page_images(page), start=1):
            figures.append({
                "data": data,
                "extension": _extension(name),
                "page": page_no,
                "order": order,
                "context": page_text[:CONTEXT_CHARS],
            })
        if len(figures) >= MAX_FIGURES:
            break
    return figures


def read_docx_figures(path: str) -> list[dict]:
    """Walk the document in order so each picture keeps the paragraphs around it."""
    import docx

    document = docx.Document(path)
    body = document.element.body
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    relationship = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

    # Paragraph text in document order, so the words near a picture can be picked up.
    paragraphs = []
    for element in body.iter():
        if element.tag.endswith("}p"):
            paragraphs.append(element)

    figures: list[dict] = []
    for position, paragraph in enumerate(paragraphs):
        blips = list(paragraph.iter(f"{namespace}blip"))
        if not blips:
            continue
        before = " ".join(
            _tidy("".join(node.text or "" for node in item.iter() if item is not None
                          and node.tag.endswith("}t")))
            for item in paragraphs[max(0, position - 4): position + 3]
        )
        for order, blip in enumerate(blips, start=1):
            rid = blip.get(f"{relationship}embed") or blip.get(f"{relationship}link")
            if not rid:
                continue
            try:
                part = document.part.related_parts[rid]
                data = part.blob
                name = getattr(part, "partname", "image.png")
            except Exception:
                continue
            figures.append({
                "data": data,
                "extension": _extension(str(name), getattr(part, "content_type", "")),
                "page": 0,
                "order": order,
                "context": _tidy(before)[:CONTEXT_CHARS],
            })
            if len(figures) >= MAX_FIGURES:
                return figures
    return figures


def read_pptx_figures(path: str) -> list[dict]:
    from pptx import Presentation

    presentation = Presentation(path)
    figures: list[dict] = []
    for index, slide in enumerate(presentation.slides, start=1):
        words = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                words.append(_tidy(shape.text_frame.text))
        context = _tidy(" ".join(words))[:CONTEXT_CHARS]
        order = 0
        for shape in slide.shapes:
            image = getattr(shape, "image", None)
            if image is None:
                continue
            try:
                data = image.blob
            except Exception:
                continue
            order += 1
            figures.append({
                "data": data,
                "extension": _extension(image.filename or "", image.content_type or ""),
                "page": index,
                "order": order,
                "context": context,
            })
            if len(figures) >= MAX_FIGURES:
                return figures
    return figures


READERS = {"pdf": read_pdf_figures, "docx": read_docx_figures, "pptx": read_pptx_figures}


def read_figures(path: str, kind: str) -> list[dict]:
    """Every usable picture in the file, in the order it appears. Never raises."""
    reader = READERS.get(kind)
    if reader is None:
        return []
    try:
        raw = reader(path)
    except Exception:
        return []

    # Count how often each picture appears before deciding what to keep. A logo on every
    # page is identical bytes each time, so the count is what gives it away.
    counts: dict[str, int] = {}
    for figure in raw:
        digest = hashlib.sha1(figure["data"]).hexdigest()
        figure["digest"] = digest
        counts[digest] = counts.get(digest, 0) + 1

    kept, seen = [], set()
    for figure in raw:
        digest = figure["digest"]
        if digest in seen or counts[digest] > MAX_REPEATS:
            continue
        width, height = _dimensions(figure["data"])
        if not _worth_keeping(figure["data"], width, height):
            continue
        if figure["extension"] not in SHOWABLE:
            continue
        seen.add(digest)
        figure["width"] = width
        figure["height"] = height
        kept.append(figure)
    return kept


# ------------------------------------------------------------------------- placing

def _overlap(context: str, text: str) -> int:
    """How much of the text around a picture turns up in a section's own text."""
    if not context or not text:
        return 0
    haystack = text.lower()
    score = 0
    for piece in re.split(r"(?<=[.!?])\s+", context):
        piece = piece.strip().lower()
        if len(piece) < 18:
            continue
        if piece[:60] in haystack:
            score += 3
    if score:
        return score
    words = [w for w in re.findall(r"[a-z]{6,}", context.lower())][:60]
    return sum(1 for word in set(words) if word in haystack)


def place_figure(figure: dict, chapters: list, modules: list) -> tuple:
    """Return (chapter, module) for this picture, matching on the text around it.

    A module is preferred, because that is the smallest section a student reads. When
    nothing matches well enough the picture goes on the chapter whose text is closest,
    and failing that on the first chapter, so no picture is quietly thrown away.
    """
    context = figure.get("context") or ""
    best_module, best_module_score = None, 0
    for module in modules:
        score = _overlap(context, module.raw_text)
        if score > best_module_score:
            best_module, best_module_score = module, score
    if best_module is not None and best_module_score >= 2:
        return best_module.chapter, best_module

    best_chapter, best_chapter_score = None, 0
    for chapter in chapters:
        score = _overlap(context, chapter.raw_text)
        if score > best_chapter_score:
            best_chapter, best_chapter_score = chapter, score
    if best_chapter is not None and best_chapter_score >= 2:
        return best_chapter, None
    return (chapters[0] if chapters else None), None


def caption_from_context(context: str, page: int) -> str:
    """A readable line to sit under the picture when no model has described it yet."""
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", context or "") if len(s.strip()) > 25]
    lead = sentences[0][:180] if sentences else ""
    where = f"Page {page} of the book" if page else "From the book"
    return f"{where}. {lead}".strip() if lead else where
