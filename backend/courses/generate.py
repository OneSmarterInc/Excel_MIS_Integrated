"""Build a book for a course when the faculty member has no book to upload.

The input is whatever they do have: lecture slides, a montage or a folder of images,
a PDF handout, or just a list of topic names. Anything with text in it is read, the
headings become the outline, and each module's prose is written by the local Ollama
model. If Ollama is not running the module is still written, from the source material
and a fixed teaching pattern, so a course can always be built.
"""
from __future__ import annotations

import json
import re
import urllib.request

from django.conf import settings

from .extraction import clean, read_docx, read_pdf, read_pptx, read_txt

IMAGE_TYPES = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff")
GENERIC = {
    "agenda", "outline", "overview", "questions", "thank you", "introduction", "contents",
    "objectives", "summary", "recap", "next week", "any questions", "title", "slide",
}
MODULES_PER_CHAPTER = 3
MAX_CHAPTERS = 6


def _looks_like_heading(text: str) -> bool:
    text = text.strip()
    if not (4 <= len(text) <= 90):
        return False
    if text.lower().strip(" .:-") in GENERIC:
        return False
    if text.endswith((",", ";")):
        return False
    return len(text.split()) <= 14


def _from_filename(name: str) -> str:
    stem = re.sub(r"\.[A-Za-z0-9]+$", "", name)
    stem = re.sub(r"[_\-]+", " ", stem)
    stem = re.sub(r"\b(img|image|photo|screenshot|montage|slide|week|final|v\d+)\b", " ",
                  stem, flags=re.IGNORECASE)
    stem = re.sub(r"\b\d{2,}\b", " ", stem)
    stem = re.sub(r"\s{2,}", " ", stem).strip()
    return stem.title() if stem else ""


def read_slides(path: str) -> dict:
    """Read a deck slide by slide, taking the biggest line on each slide as its heading.

    Decks built from text boxes have no title placeholder, so the type size is what marks
    the heading. Small all caps lines above it are kickers, and the repeated line along the
    bottom is a footer; neither is a topic.
    """
    from pptx import Presentation

    presentation = Presentation(path)
    headings, lines, footers = [], [], {}

    slides = []
    for slide in presentation.slides:
        entries = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = clean(" ".join(run.text for run in paragraph.runs))
                if not text:
                    continue
                sizes = [run.font.size.pt for run in paragraph.runs if run.font.size]
                entries.append({"text": text, "size": max(sizes) if sizes else 0,
                                "top": shape.top or 0})
                footers[text] = footers.get(text, 0) + 1
        slides.append(entries)

    repeated = {text for text, count in footers.items() if count >= max(3, len(slides) // 2)}

    for entries in slides:
        usable = [e for e in entries
                  if e["text"] not in repeated and not e["text"].strip().isdigit()]
        if not usable:
            continue
        biggest = max(usable, key=lambda e: e["size"])
        heading = biggest["text"].replace("\n", " ")
        if biggest["size"] >= 18 and _looks_like_heading(heading):
            headings.append(heading)
        for entry in usable:
            if entry is biggest:
                continue
            if len(entry["text"]) > 25 and entry["size"] < biggest["size"]:
                lines.append(entry["text"])
    return {"headings": headings, "lines": lines}


def read_source(path: str, filename: str) -> dict:
    """Return headings and supporting lines from one uploaded file."""
    lowered = filename.lower()
    headings: list[str] = []
    lines: list[str] = []

    if lowered.endswith((".pptx", ".ppt")):
        slides = read_slides(path)
        if slides["headings"]:
            return slides
        for block in read_pptx(path):
            if block["style"] == "h2" and _looks_like_heading(block["text"]):
                headings.append(block["text"])
            elif len(block["text"]) > 20:
                lines.append(block["text"])
    elif lowered.endswith(".pdf"):
        for block in read_pdf(path):
            if _looks_like_heading(block["text"]) and block["text"][:1].isupper():
                headings.append(block["text"])
            elif len(block["text"]) > 30:
                lines.append(block["text"])
    elif lowered.endswith((".docx", ".doc")):
        for block in read_docx(path):
            if block["style"] in ("h1", "h2") or _looks_like_heading(block["text"]):
                headings.append(block["text"])
            elif len(block["text"]) > 30:
                lines.append(block["text"])
    elif lowered.endswith(IMAGE_TYPES):
        # An image carries no readable text, so its file name stands in as a topic. That is a
        # weak signal, used only when the slides and documents gave us little to go on.
        guess = _from_filename(filename)
        return {"headings": [], "lines": [], "weak_headings": [guess] if guess else []}
    else:
        for block in read_txt(path):
            if _looks_like_heading(block["text"]):
                headings.append(block["text"])
            elif len(block["text"]) > 30:
                lines.append(block["text"])

    return {"headings": headings, "lines": lines}


def gather(sources: list[dict], topics: str = "") -> dict:
    headings: list[str] = []
    weak: list[str] = []
    lines: list[str] = []
    for source in sources:
        headings += source.get("headings", [])
        weak += source.get("weak_headings", [])
        lines += source.get("lines", [])
    if len(headings) < 6:
        headings += weak

    # Teachers type a topic list any way they like: one per line, or all on one line
    # separated by commas or semicolons. A comma inside a single topic is rare enough
    # that splitting on it is the right trade.
    typed: list[str] = []
    for raw in re.split(r"[\n;,]+", topics or ""):
        topic = clean(raw).strip(" -•*.")
        if len(topic) >= 4:
            typed.append(topic)

    # A montage carries no text, so its file name is only a heading of last resort. What
    # the teacher typed comes first, then anything read out of a deck or a document.
    headings = typed + headings

    seen, unique = set(), []
    for heading in headings:
        key = re.sub(r"[^a-z0-9]+", "", heading.lower())
        if key and key not in seen:
            seen.add(key)
            unique.append(clean(heading).strip(" -:.")[:120])
    return {"headings": unique, "lines": lines}


def usable_title(text: str) -> bool:
    """Reject the fragments that slip out of a PDF and read badly as a chapter title."""
    cleaned = (text or "").strip()
    if not 8 <= len(cleaned) <= 90:
        return False
    words = cleaned.split()
    if not 2 <= len(words) <= 12:
        return False
    if cleaned[0] in "),.;:-" or cleaned.endswith((",", ";", "(")):
        return False
    if cleaned.count("(") != cleaned.count(")"):
        return False
    if re.search(r"[=<>]|\$[A-Z]|\bB\d|\bC\d", cleaned):   # a formula, not a heading
        return False
    if sum(character.isdigit() for character in cleaned) > len(cleaned) / 3:
        return False
    return any(character.isalpha() for character in cleaned)


NUMBERED_HEADING = re.compile(
    r"^\s*(?:chapter|module|section|unit|part|lesson|topic)\s+\d{1,2}(?:\.\d{1,2})?\s*[:.\-\u2013]?\s*(.{3,90})$",
    re.IGNORECASE,
)


def preferred_headings(headings: list[str]) -> list[str]:
    """When a document numbers its own headings, those are the real ones.

    A PDF gives up dozens of short lines, and most of them are table cells or the first
    line of a paragraph. If the document says "Chapter 2: Formulas and Cell References"
    anywhere, trust that and ignore the rest.
    """
    numbered = []
    for heading in headings:
        match = NUMBERED_HEADING.match(heading)
        if match:
            title = match.group(1).strip(" .:-")
            if len(title) >= 4:
                numbered.append(title)
    return numbered if len(numbered) >= 4 else []


def outline(course_name: str, material: dict) -> list[dict]:
    """Group the headings into chapters of three modules each."""
    headings = preferred_headings(material["headings"])
    headings = headings or [head for head in material["headings"] if usable_title(head)]
    headings = headings or material["headings"] or [course_name or "Course overview"]
    while len(headings) < MODULES_PER_CHAPTER * 2:
        headings.append(f"{headings[len(headings) % len(headings)]} in practice")

    # Each group gives one chapter: the first heading names the chapter, the rest are its
    # modules, so a title never appears twice on the same page.
    per_group = MODULES_PER_CHAPTER + 1
    chapter_count = max(2, min(MAX_CHAPTERS, max(1, round(len(headings) / per_group))))
    per_group = max(2, -(-len(headings) // chapter_count))

    chapters = []
    for index in range(chapter_count):
        start = index * per_group
        group = headings[start:] if index == chapter_count - 1 else headings[start:start + per_group]
        if not group:
            continue
        title = group[0]
        module_titles = group[1:per_group] or [f"{title} in practice"]
        chapters.append({
            "number": len(chapters) + 1,
            "title": title[:120],
            "modules": [
                {"number": position, "title": module_title[:120]}
                for position, module_title in enumerate(module_titles[:MODULES_PER_CHAPTER], start=1)
            ],
        })
    return chapters


# --------------------------------------------------------------------------- writing

def _context_for(title: str, lines: list[str], limit: int = 6) -> list[str]:
    words = {w for w in re.findall(r"[a-z]{4,}", title.lower())}
    scored = []
    for line in lines:
        hits = len(words & set(re.findall(r"[a-z]{4,}", line.lower())))
        if hits:
            scored.append((hits, line))
    scored.sort(key=lambda item: item[0], reverse=True)
    return [line for _, line in scored[:limit]]


def ollama_write(course_name: str, chapter_title: str, module_title: str,
                 context: list[str]) -> str | None:
    if not getattr(settings, "USE_OLLAMA", True):
        return None
    host = getattr(settings, "OLLAMA_HOST", "http://127.0.0.1:11434").rstrip("/")
    model = getattr(settings, "OLLAMA_MODEL", "llama3.1:8b")
    timeout = getattr(settings, "OLLAMA_TIMEOUT", 180)

    prompt = (
        f"You are writing one module of a textbook for the course '{course_name}'.\n"
        f"The chapter is '{chapter_title}' and this module is '{module_title}'.\n\n"
        "Write about 400 words of plain English teaching prose for undergraduate business "
        "students. Explain the idea, say why it matters in a real organisation, and give one "
        "short concrete example. Do not use markdown symbols, bullet characters or headings. "
        "Do not mention slides, lectures or this instruction.\n\n"
        + ("Material from the lecture to stay close to:\n" + "\n".join(context) if context else "")
    )
    body = json.dumps({
        "model": model,
        "stream": False,
        "options": {"temperature": 0.4, "num_predict": 1200},
        "messages": [{"role": "user", "content": prompt}],
    }).encode()
    request = urllib.request.Request(
        f"{host}/api/chat", data=body, headers={"content-type": "application/json"}
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode())
    except Exception:
        return None
    text = (payload.get("message") or {}).get("content", "").strip()
    if len(text) < 200:
        return None
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    return text.replace("**", "").replace("*", "").replace("`", "").strip()


def local_write(course_name: str, chapter_title: str, module_title: str,
                context: list[str]) -> str:
    """Written from the material itself when the model is not available.

    It follows the same shape a lecturer would use: what the idea is, what the source
    material says about it, a worked step, what goes wrong, and a check. The lecturer's own
    lines carry most of the weight, so the result stays on the topic rather than drifting
    into generalities.
    """
    subject = module_title.lower().strip(" .:")
    body = [
        f"{module_title} sits inside {chapter_title.lower()} and is part of {course_name}. "
        "This module sets out what the idea is, where it shows up in an organisation, and "
        "what you are expected to be able to do with it by the end of the session."
    ]

    if context:
        body.append("What the material says\n" + " ".join(context[:6]))
        if len(context) > 6:
            body.append(" ".join(context[6:10]))
    else:
        body.append(
            f"What the material says\nThe slides for this session name {subject} without "
            "spelling it out, so write your own one sentence definition of it during the "
            "class and keep that sentence beside your notes."
        )

    body.append(
        f"Working through it\nTake the smallest example of {subject} you can build, put it "
        "in a worksheet with your own figures, and change one input at a time. Watch which "
        "results move and which stay still. That tells you what the idea depends on far "
        "faster than reading it twice does."
    )
    body.append(
        f"Where it goes wrong\nMost mistakes with {subject} come from applying it without "
        "checking what it assumes. Before you use a result, ask what the numbers behind it "
        "would have to be true for, and whether anybody has checked that they are."
    )
    body.append(
        f"In practice, {subject} matters because decisions get made on the back of it. A "
        "manager reading your work will not rebuild it, so what you set out here is what "
        "they act on."
    )
    body.append(
        f"Check yourself\nCan you explain {subject} to somebody who missed the class, show "
        "it working, and say what a manager would do differently because of it?"
    )
    return "\n\n".join(body)


def write_module(course_name: str, chapter_title: str, module_title: str,
                 lines: list[str]) -> str:
    context = _context_for(module_title, lines)
    return (ollama_write(course_name, chapter_title, module_title, context)
            or local_write(course_name, chapter_title, module_title, context))


def build_book(course_name: str, material: dict) -> dict:
    """Return the finished chapter and module structure, ready to be stored."""
    plan = outline(course_name, material)
    lines = material["lines"]
    chapters = []
    for chapter in plan:
        modules = []
        for module in chapter["modules"]:
            text = write_module(course_name, chapter["title"], module["title"], lines)
            modules.append({
                "number": module["number"],
                "title": module["title"],
                "raw_text": f"Module {module['number']}: {module['title']}\n\n{text}",
            })
        chapter_text = "\n\n".join(
            [f"Chapter {chapter['number']}: {chapter['title']}"] +
            [module["raw_text"] for module in modules]
        )
        chapters.append({
            "number": chapter["number"],
            "title": chapter["title"],
            "raw_text": chapter_text,
            "modules": modules,
        })
    full_text = "\n\n".join(chapter["raw_text"] for chapter in chapters)
    return {
        "chapters": chapters,
        "text": f"{course_name}\n\n{full_text}",
        "character_count": len(full_text),
    }
